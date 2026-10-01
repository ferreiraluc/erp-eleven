"""Individual-row BI tests use only synthetic values and an isolated database."""
from copy import deepcopy
from datetime import date, datetime, time, timezone
from io import BytesIO
import json
from types import SimpleNamespace

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.api.endpoints import sales_bi_entries
from app.database import get_db
from app.dependencies import get_current_user
from app.models.sales_bi import SalesBIWorkbook
from app.services.sales_bi_entry_parser import extract_entries, observed_date, observed_time
from app.services.sales_bi_entries import build_entries
from app.services.sales_bi_parser import parse_workbook, PARSER_VERSION
from test_sales_bi import fixture_workbook


def detail_sheet():
    rows = [[None] * 60 for _ in range(25)]
    rows[10][1:8] = ['Moeda', 'Valor Bruto', 'Vendedor', 'Metodo Pgto', 'Valor Liquido', 'Hora', 'Cliente']
    rows[11][0] = date(2026, 9, 23)
    rows[12][1:8] = ['R$', 100, 'Junior', 'PIX', 98, time(10, 30), 'Cliente de teste']
    rows[13][1:8] = ['R$', 100, 'Juninho', 'PIX', 98, time(10, 30), 'Cliente de teste']
    rows[14][1:8] = ['U$', 50, 'Lucas', 'Dinheiro', 50, '11:45', None]
    rows[15][0] = 'TER'
    rows[16][1:8] = ['G$', 7000, 'Sol', None, 7000, '12:00', 'Exemplo']
    rows[17][1:4] = ['R$', None, 'Junior']
    rows[18][1:4] = ['R$', 300, None]
    rows[19][1:4] = ['R$', '=SUM(A1)', 'Junior']
    rows[20][0] = 'observação sem data'
    rows[21][1:8] = ['R$', 80, 'Junior', None, None, None, None]
    return {'rows': rows, 'total_row': 8}


def details():
    return extract_entries([('semana1', detail_sheet())], [{'sheet': 'semana1', 'index': 1, 'label': '21/09 - 27/09'}], year=2026)


def workbook(key='one', kind='archive', *, snapshot=None):
    return SalesBIWorkbook(id=key, filename='Setembro.xlsx', kind=kind, active=True, year=2026, month=9,
                           remote_version='1', synced_at=datetime(2026, 9, 30, tzinfo=timezone.utc),
                           snapshot=snapshot or {'total_usd': '777', 'sellers': {
                               'Junior': {'total_usd': '100', 'currencies': {'BRL': '196'}},
                               'Lucas': {'total_usd': '50', 'currencies': {'USD': '50'}},
                               'Sol': {'total_usd': '1', 'currencies': {'PYG': '7000'}}}, 'entries': details()})


def test_entries_preserve_equal_real_rows_and_canonical_sellers():
    data = details()
    rows = data['rows']
    assert len(rows) == 5
    assert [r['row'] for r in rows] == [13, 14, 15, 17, 22]
    assert rows[0]['seller'] == rows[1]['seller'] == 'Junior'
    assert rows[0]['gross'] == rows[1]['gross'] == '100'
    assert rows[0]['fingerprint'] != rows[1]['fingerprint']
    assert rows[0]['source_cell'] == 'semana1!C13'
    assert rows[0]['net'] == '98'
    assert rows[-1]['net'] is None
    assert data['diagnostics'] == [{'sheet': 'semana1', 'skipped_rows': 3}]
    assert '=SUM' not in json.dumps(data)


def test_intraday_requires_real_date_and_time_not_week_label_or_sync():
    rows = details()['rows']
    assert rows[0]['date'] == '2026-09-23' and rows[0]['time'] == '10:30:00'
    assert rows[2]['time'] == '11:45:00'
    assert rows[3]['date'] is None and rows[3]['time'] is None and rows[3]['day_group'] == 'tue'
    assert rows[4]['date'] is None
    result = build_entries([workbook()])
    assert result['summary']['count'] == 5
    assert result['summary']['dated_count'] == 3
    assert result['summary']['timed_count'] == 3
    assert [(v['hour'], v['count']) for v in result['hourly']] == [(10, 2), (11, 1)]
    assert result['dates'] == ['2026-09-23']
    assert result['weekdays'][0]['day'] == 'tue' and result['weekdays'][0]['count'] == 1
    assert result['summary']['official_total_usd'] == 777


def test_optional_columns_need_headers_and_legacy_net_uses_gross():
    data = detail_sheet()
    data['rows'][10] = [None] * 60
    for row in data['rows']:
        row[5] = None
    entry = extract_entries([('semana1', data)], [], year=2026)['rows'][0]
    assert entry['net'] == '100'
    assert entry['customer'] is None and entry['time'] is None
    assert entry['week_index'] is None


@pytest.mark.parametrize('value,expected', [
    (datetime(2026, 9, 23, 16, 20), ('2026-09-23', '16:20:00')),
    (datetime(2026, 9, 23), ('2026-09-23', None)),
    ('23/09/2026 16:20', ('2026-09-23', '16:20:00')),
    ('23/09', ('2026-09-23', None)), ('31/09', (None, None)),
    ('23/09 - 29/09', (None, None)), ('SEG', (None, None)),
])
def test_dates_are_explicit_and_invalid_dates_do_not_invent_values(value, expected):
    assert observed_date(value, 2026) == expected


def test_explicit_midnight_time_is_kept_and_excel_serial_is_not_guessed():
    assert observed_time('00:00') == '00:00:00'
    assert observed_time(time(0)) == '00:00:00'
    assert observed_time(.5) is None
    assert observed_time('25:05') is None


def test_partial_date_in_different_month_does_not_guess_year_rollover():
    assert observed_date('31/12', 2026, month=1) == (None, None)
    assert observed_date('31/12/2025', 2026, month=1) == ('2025-12-31', None)


def test_combined_timestamp_column_can_explicitly_record_midnight():
    data = detail_sheet()
    data['rows'][10][0] = 'Data e hora'
    data['rows'][12][0] = datetime(2026, 9, 23, 0, 0)
    data['rows'][12][6] = None
    entry = extract_entries([('semana1', data)], [], year=2026, month=9)['rows'][0]
    assert entry['date'] == '2026-09-23' and entry['time'] == '00:00:00'


def test_workbook_snapshot_excludes_only_full_live_mirror():
    content = fixture_workbook(duplicate=True, monthly=False)
    data = parse_workbook(content, current=True)
    assert data['version'] == PARSER_VERSION == 3
    assert {r['sheet'] for r in data['entries']['rows']} == {'semana 1', 'semana2'}
    assert data['total_usd'] == '1000'
    # Equal payments on different recorded dates do not establish a copied tab.
    book = load_workbook(BytesIO(content), data_only=True)
    book['Planilha1']['A12'] = date(2026, 9, 23)
    book['semana2']['A12'] = date(2026, 9, 22)
    output = BytesIO(); book.save(output); book.close()
    parsed = parse_workbook(output.getvalue(), current=True)
    assert {r['sheet'] for r in parsed['entries']['rows']} == {'Planilha1', 'semana 1', 'semana2'}


def test_current_archive_precedence_and_pagination_use_same_scoped_rows():
    current = workbook('current', 'current')
    archive = workbook()
    result = build_entries([current, archive], seller='Junior', offset=1, limit=1, private=True)
    assert result['total'] == 3
    assert len(result['items']) == 1
    assert result['items'][0]['source_id'] == 'one'
    assert result['sellers'] == ['Junior']
    assert result['summary']['official_total_usd'] == 100
    assert result['summary']['currencies'] == [{'currency': 'BRL', 'count': 3, 'gross': 280, 'net': 196, 'net_available': 2}]
    assert result['coverage']['skipped_rows'] is None
    assert 'Lucas' not in json.dumps(result, default=str) and 'Sol' not in json.dumps(result, default=str)


def test_day_currency_search_apply_to_table_and_aggregates_without_changing_closing():
    result = build_entries([workbook()], seller='Junior', currency='BRL', day='2026-09-23', search='teste')
    assert result['total'] == result['summary']['count'] == 2
    assert result['summary']['currencies'][0]['gross'] == 200
    assert result['hourly'][0]['count'] == 2
    assert result['summary']['official_total_usd'] == 100


def test_old_snapshot_requires_sync_without_fabricating_zero_detail():
    source = workbook()
    source.snapshot = deepcopy(source.snapshot); source.snapshot.pop('entries')
    result = build_entries([source])
    assert result['items'] == [] and result['coverage']['needs_sync']
    assert result['coverage']['sources_without_entries'] == 1
    assert result['summary']['official_total_usd'] == 777


@pytest.fixture
def client(tmp_path):
    engine = create_engine('sqlite:///' + str(tmp_path/'entries.db'), connect_args={'check_same_thread': False})
    SalesBIWorkbook.__table__.create(engine)
    factory = sessionmaker(bind=engine)
    with factory() as db:
        db.add(workbook()); db.commit()
    app = FastAPI(); app.include_router(sales_bi_entries.router, prefix='/bi')
    def database():
        with factory() as db:
            yield db
    app.dependency_overrides[get_db] = database
    user = SimpleNamespace(role=SimpleNamespace(value='GERENTE'), ativo=True, sales_scope='own', sales_seller='Junior')
    app.dependency_overrides[get_current_user] = lambda: user
    with TestClient(app) as browser:
        yield browser, user
    engine.dispose()


def test_endpoint_overrides_forged_seller_and_filters_every_aggregate(client):
    browser, _ = client
    response = browser.get('/bi/entries?seller=Lucas')
    assert response.status_code == 200
    body = response.json()
    assert body['access'] == {'scope': 'own', 'seller': 'Junior'}
    assert body['summary']['official_total_usd'] == 100
    assert body['summary']['count'] == 3
    assert 'Lucas' not in response.text and 'Sol' not in response.text
    assert browser.get('/bi/entries?currency=USD').json()['items'] == []


def test_endpoint_fails_closed_without_seller_binding_and_validates_filters(client):
    browser, user = client
    user.sales_seller = None
    assert browser.get('/bi/entries').status_code == 403
    user.sales_seller = 'Junior'
    assert browser.get('/bi/entries?day=2026-09-31').status_code == 422
    assert browser.get('/bi/entries?limit=1000').status_code == 422
    assert browser.get('/bi/entries?currency=ZZZ').status_code == 422
    user.role = SimpleNamespace(value='VENDEDOR')
    assert browser.get('/bi/entries').status_code == 403


def test_global_access_has_all_rows_without_mutating_or_creating_sales(client):
    browser, user = client
    user.sales_scope = 'all'
    result = browser.get('/bi/entries').json()
    assert result['total'] == 5
    assert result['sellers'] == ['Junior', 'Lucas', 'Sol']
    assert result['summary']['official_total_usd'] == 777
    assert result['coverage']['skipped_rows'] == 3
