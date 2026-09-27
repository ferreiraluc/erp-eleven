"""Isolated BI tests: generated workbooks, SQLite and stubbed GET sources only."""
import copy
import importlib.util
from datetime import timedelta, datetime
from io import BytesIO
import json
import os
from pathlib import Path
from types import SimpleNamespace
from xml.etree import ElementTree as ET
from zipfile import ZipFile

os.environ['DATABASE_URL'] = 'sqlite://'
import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from fastapi import FastAPI
from fastapi.testclient import TestClient
from app.database import Base, get_db
from app.dependencies import get_current_user
from app.models.sales_bi import SalesBIConfig, SalesBIWorkbook
from app.models.assistant import utcnow
from app.services.sales_bi_parser import parse_workbook, WorkbookError, filename_month
from app.services.sales_bi_onedrive import OneDriveReader, SourceError, validate_url, validate_root
from app.services.sales_bi import build_overview, sources_status
from app.services.sales_bi_sync import sync_once, claim
from app.api.endpoints import sales_bi

CONFIG = {'current_url': 'https://1drv.ms/x/example', 'archive_url': 'https://1drv.ms/f/example',
          'archive_root': '/personal/0123456789abcdef/Documents/History'}
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


def fixture_workbook(*, old=False, duplicate=False, missing_cache=False, named='Planilha1', mismatch=False, monthly=True, date_label=False, shifted=False, currency_text=False):
    w = Workbook(); w.active.title = named
    tr = 7 if old else 8
    for name, junior, lucas in [(named, 100, 450), ('semana 1', 200, 250)]:
        s = w[name] if name in w.sheetnames else w.create_sheet(name)
        s['D2'], s['H2'] = 'Junior', 'Lucas'
        s['C3'], s['C4'], s['C5'] = 'G$', 'R$', 'U$'
        s['D3'], s['H3'], s['D4'], s['H4'], s['D5'], s['H5'] = 7000, 14000, 50, 100, junior, lucas
        s[f'B{tr}'] = 'total';s[f'D{tr}'] = junior;s[f'H{tr}'] = lucas;s[f'K{tr}'] = junior+lucas
        s['B12'], s['C12'], s['D12'], s['F12'] = 'U$', lucas, 'Lucas', lucas
    s = w[named]
    s['R2'], s['S2'] = datetime(2026,9,1) if date_label else '01/09 - 07/09', 450
    s['R3'], s['S3'] = '08/09 - 14/09', 550
    s['R10'], s['S10'], s['T12'] = 'Total:', '=SUM(S2:S3)+200-300', 2026
    if monthly:
        s['W5'], s['W6'], s['X6'], s['W7'], s['X7'] = 'Total Mês', 'Juninho', 'Lucas', 300, 700
    if duplicate:
        w.copy_worksheet(s).title = 'semana2'
    if mismatch:
        w['semana 1'][f'K{tr}'] = 999
    if currency_text:
        w['semana 1'][f'D{tr}'] = '$200,00'
        w['semana 1'][f'H{tr}'] = '$250,00'
        w['semana 1'][f'K{tr}'] = '$450,00'
    s[f'H{tr}'] = '=SUM(H5:H6)+200-300'
    output=BytesIO();w.save(output)
    result=BytesIO()
    with ZipFile(BytesIO(output.getvalue())) as source, ZipFile(result,'w') as target:
        for name in source.namelist():
            content = source.read(name)
            if name.startswith('xl/worksheets/sheet'):
                root=ET.fromstring(content)
                for c in root.findall('.//s:c',NS):
                    formula=c.find('s:f',NS)
                    if formula is not None and not missing_cache:
                        v=c.find('s:v',NS)
                        if v is None:v=ET.SubElement(c,'{'+NS['s']+'}v')
                        v.text='1000' if c.attrib['r']=='S10' else '450'
                if shifted and name=='xl/worksheets/sheet2.xml':
                    for row in root.findall('.//s:row',NS):
                        if int(row.attrib['r'])<10:
                            row.attrib['r']=str(int(row.attrib['r'])+1)
                            for c in row.findall('s:c',NS):
                                import re
                                col,r=re.fullmatch(r'([A-Z]+)([0-9]+)',c.attrib['r']).groups()
                                c.attrib['r']=col+str(int(r)+1)
                content=ET.tostring(root)
            target.writestr(name,content)
    return result.getvalue()


def snapshot(**kwargs):
    return parse_workbook(fixture_workbook(**kwargs),current=False)


def row(key='a', *, kind='current', year=2026, month=9, data=None, version='1', error=None):
    return SalesBIWorkbook(id=key,kind=kind,filename=key+'.xlsx',year=year,month=month,
                           snapshot=data or snapshot(),remote_version=version,active=True,error=error,synced_at=utcnow())


@pytest.fixture
def factory(tmp_path):
    engine=create_engine('sqlite:///'+str(tmp_path/'bi.db'),connect_args={'check_same_thread':False})
    for table in (SalesBIConfig.__table__,SalesBIWorkbook.__table__):table.create(engine)
    yield sessionmaker(bind=engine)
    engine.dispose()


def test_saved_corrected_results_not_formula_recalculation():
    data=snapshot()
    assert data['total_usd']=='1000'
    assert data['sellers']['Lucas']['total_usd']=='700'
    assert data['weeks'][1]['sellers']['Lucas']=='450'
    assert data['sellers']['Lucas']['currencies']['BRL']=='200'
    assert data['year']==2026 and data['month']==9
    assert '+200' not in json.dumps(data) and 'SUM' not in json.dumps(data)


def test_missing_excel_cache_never_becomes_zero():
    with pytest.raises(WorkbookError,match='resultado salvo'):
        snapshot(missing_cache=True)


def test_duplicate_current_week_is_not_counted_twice():
    data=parse_workbook(fixture_workbook(duplicate=True,monthly=False),current=True)
    assert len(data['weeks'])==2
    assert data['sellers']['Lucas']['total_usd']=='700'
    assert data['weeks'][1]['sheet']=='semana2'


@pytest.mark.parametrize('named',['Planilha1','Vendas','Plhanilha1'])
def test_historical_layout_names_and_currency_row_shift(named):
    data=snapshot(old=True,named=named)
    assert data['sellers']['Junior']['total_usd']=='300'
    assert all(w['detail_available'] for w in data['weeks'])
    assert data['weeks'][0]['sheet']=='semana 1'


def test_unmatched_week_omits_derived_detail_preserves_published_monthly_seller():
    data=snapshot(mismatch=True)
    assert data['total_usd']=='1000'
    assert data['weeks'][0]['detail_available'] is False
    assert data['sellers']['Lucas']['total_usd']=='700'
    assert data['sellers']['Lucas']['currencies']=={}
    assert data['warnings']
    assert snapshot(mismatch=True,monthly=False)['sellers']['Lucas']['total_usd'] is None


def test_historical_period_uses_folder_year_and_filename_not_worksheet_template():
    data=parse_workbook(fixture_workbook(),year=2024,month=2)
    assert (data['year'],data['month'])==(2024,2)
    assert filename_month('Maio-Fatuarmento.xlsx')==5
    assert filename_month('Março-Faturamento.xlsx')==3
    assert filename_month('cópia.xlsx') is None


def test_archive_and_current_period_are_counted_once_and_ranked():
    a=snapshot();b=copy.deepcopy(a);b['total_usd']='2000'
    result=build_overview([row(),row('archive',kind='archive',data=b),row('older',kind='archive',year=2025)],year=2026,month=9)
    assert result['selected']['total_usd']==2000
    assert result['selected']['source_months']==1
    assert result['selected']['partial'] is False
    assert result['comparison'][1]['difference_usd']==1000
    assert result['annual'][0]['total_usd']==2000
    assert result['ranking'][0]['name']=='Lucas'
    assert len(result['weeks'])==2
    assert result['currencies'][1]['value']==300


def test_absence_is_not_zero_and_rankings_disclose_coverage():
    d=snapshot(mismatch=True,monthly=False)
    result=build_overview([row(),row('old',kind='archive',year=2025,data=d)],seller='Lucas')
    assert result['selected']['total_usd']==700
    assert result['selected']['available_months']==1
    assert result['selected']['source_months']==2
    assert result['comparison'][0]['total_usd'] is None
    assert result['comparison'][1]['difference_percent'] is None
    assert build_overview([],year=2025)['selected']['total_usd'] is None


class Reader:
    downloads=0
    fail=False
    def __init__(self,*args):pass
    def current(self):return fixture_workbook()
    def archive(self):return [{'remote_id':'a','filename':'Setembro.xlsx','year':2024,'month':9,'path':'file','version':'1'}]
    def download(self,path):
        type(self).downloads+=1
        if self.fail:raise SourceError('Falha de leitura de teste.')
        return fixture_workbook()
    def close(self):pass


def configured(factory):
    with factory() as db:
        db.add(SalesBIConfig(id=1,**CONFIG));db.commit()


def due(factory):
    with factory() as db:
        c=db.get(SalesBIConfig,1);c.next_sync_at=utcnow()-timedelta(minutes=1);db.commit()


def test_sync_is_idempotent_skips_unchanged_and_preserves_last_good(factory):
    Reader.downloads=0;Reader.fail=False
    configured(factory)
    assert sync_once(factory,Reader)
    assert not sync_once(factory,Reader)
    due(factory);assert sync_once(factory,Reader)
    assert Reader.downloads==1
    with factory() as db:
        rows=db.query(SalesBIWorkbook).all();assert len(rows)==2
        archive=next(r for r in rows if r.kind=='archive');archive.remote_version='old';db.commit()
    Reader.fail=True;due(factory);assert sync_once(factory,Reader)
    with factory() as db:
        archive=db.query(SalesBIWorkbook).filter_by(kind='archive').one()
        assert archive.snapshot['total_usd']=='1000' and archive.error
        assert db.get(SalesBIConfig,1).last_error
        assert not sources_status(db)['running']
    Reader.fail=False


def test_failed_listing_does_not_retire_or_delete_history(factory):
    configured(factory);Reader.fail=False;sync_once(factory,Reader)
    class Broken(Reader):
        def archive(self):raise SourceError('Histórico temporariamente indisponível.')
    due(factory);sync_once(factory,Broken)
    with factory() as db:
        assert db.query(SalesBIWorkbook).filter_by(active=True).count()==2
        assert db.get(SalesBIConfig,1).last_error


def test_database_lease_prevents_overlapping_sync_and_recovers_expiration(factory):
    configured(factory)
    assert claim(factory)
    assert claim(factory) is None
    with factory() as db:
        c=db.get(SalesBIConfig,1);c.lease_until=utcnow()-timedelta(seconds=1);db.commit()
    assert claim(factory)


@pytest.mark.parametrize('url',['http://1drv.ms/test','https://evil.example/test','https://onedrive.live.com.evil.example/test','https://user:pass@onedrive.live.com/test','https://127.0.0.1/'])
def test_connector_rejects_unapproved_destinations(url):
    with pytest.raises(SourceError):validate_url(url)


def test_connector_stays_in_shared_folder_and_uses_only_get(monkeypatch):
    reader=OneDriveReader(**CONFIG)
    with pytest.raises(SourceError):reader.download('/personal/0123456789abcdef/Documents/Other/file.xlsx')
    calls=[]
    class Response:
        status_code=302
        headers={'Location':'https://127.0.0.1/private'}
        def __enter__(self):return self
        def __exit__(self,*args):pass
    def get(url,**kwargs):calls.append(url);return Response()
    monkeypatch.setattr(reader.session,'get',get)
    with pytest.raises(SourceError):reader.current()
    assert len(calls)==1
    with pytest.raises(SourceError):validate_root(CONFIG['archive_root']+'/../Other')
    with pytest.raises(SourceError):reader.download(CONFIG['archive_root']+'/%2e%2e/Other.xlsx')
    reader.close()


@pytest.fixture
def client(factory):
    app=FastAPI();app.include_router(sales_bi.router,prefix='/bi')
    def database():
        with factory() as db:yield db
    app.dependency_overrides[get_db]=database
    def user():return SimpleNamespace(role=SimpleNamespace(value='ADMIN'),ativo=True)
    app.dependency_overrides[get_current_user]=user
    with TestClient(app) as c:yield c,app


def test_financial_data_requires_manager_and_config_requires_admin(client):
    c,app=client
    app.dependency_overrides[get_current_user]=lambda:SimpleNamespace(role=SimpleNamespace(value='VENDEDOR'),ativo=True)
    assert c.get('/bi/overview').status_code==403
    assert c.get('/bi/sources').status_code==403
    assert c.post('/bi/sync').status_code==403
    app.dependency_overrides[get_current_user]=lambda:SimpleNamespace(role=SimpleNamespace(value='GERENTE'),ativo=True)
    assert c.get('/bi/overview').status_code==200
    assert c.get('/bi/config').status_code==403
    assert c.put('/bi/config',json=CONFIG).status_code==403


def test_configuration_queue_validation_and_no_secrets_in_status(client,factory):
    c,_=client
    assert c.post('/bi/sync').status_code==400
    assert c.put('/bi/config',json=CONFIG|{'current_year':2026}).status_code==422
    assert c.put('/bi/config',json=CONFIG).status_code==200
    assert c.post('/bi/sync').status_code==202
    assert 'current_url' not in c.get('/bi/sources').json()
    assert c.get('/bi/overview?month=13').status_code==422
    claim(factory)
    assert c.put('/bi/config',json=CONFIG).status_code==409


def test_migration_creates_only_bi_tables_and_roundtrips(tmp_path):
    from alembic.migration import MigrationContext
    from alembic.operations import Operations
    from sqlalchemy import inspect
    path=Path(__file__).parents[1]/'alembic/versions/x4y5z6a7b8c9_sales_bi.py'
    spec=importlib.util.spec_from_file_location('bi_migration',path)
    module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
    engine=create_engine('sqlite://')
    with engine.begin() as conn:
        module.op=Operations(MigrationContext.configure(conn))
        module.upgrade();assert set(inspect(conn).get_table_names())=={'sales_bi_config','sales_bi_workbooks'}
        module.downgrade();assert inspect(conn).get_table_names()==[]


def test_live_month_uses_saved_week_results_even_if_manual_monthly_recap_is_stale():
    data=parse_workbook(fixture_workbook(mismatch=True),current=True)
    assert data['total_usd']=='1549'
    assert data['sellers']['Lucas']['total_usd']=='700'
    assert data['weeks'][0]['total_usd']=='999'
    assert data['source_cell'] is None


def test_current_workbook_resolves_only_the_shared_document_guid(monkeypatch):
    reader=OneDriveReader(**CONFIG)
    calls=[]
    content=fixture_workbook()
    class Response:
        status_code=200
        headers={}
        url='https://onedrive.live.com/personal/0123456789abcdef/_layouts/15/Doc.aspx?sourcedoc=%7B11111111-2222-3333-4444-555555555555%7D'
        def __enter__(self):return self
        def __exit__(self,*args):pass
        def iter_content(self,size):yield b'<html>Shared document</html>' if len(calls)==1 else content
    def get(url,**kwargs):calls.append(url);return Response()
    monkeypatch.setattr(reader.session,'get',get)
    assert reader.current()==content
    assert len(calls)==2
    assert calls[1].endswith("/_api/web/GetFileById('11111111-2222-3333-4444-555555555555')/$value")
    reader.close()


def test_excel_date_week_labels_remain_in_week_ranking_and_currency_breakdown():
    data=snapshot(date_label=True)
    assert len(data['weeks'])==2
    assert data['weeks'][0]['label']=='01/09'
    assert data['weeks'][0]['total_usd']=='450'
    assert data['currencies_complete']


def test_inserted_row_in_week_locates_header_currency_and_total_together():
    data=snapshot(shifted=True,monthly=False)
    assert data['weeks'][0]['sheet']=='semana 1'
    assert data['weeks'][0]['sellers']['Lucas']=='250'
    assert data['sellers']['Lucas']['total_usd']=='700'
    assert data['currencies_complete']


def test_manual_general_adjustment_is_preserved_and_disclosed_not_reallocated():
    data=parse_workbook(fixture_workbook(mismatch=True),current=True)
    assert data['total_usd']=='1549'
    assert data['sellers']['Lucas']['total_usd']=='700'
    assert any('549,00' in w and 'abaixo' in w for w in data['warnings'])


def test_saved_usd_text_still_reconciles_week_and_individual_seller_results():
    data=snapshot(currency_text=True,monthly=False)
    assert data['weeks'][0]['detail_available']
    assert data['weeks'][0]['sellers']['Lucas']=='250.00'
    assert data['sellers']['Lucas']['total_usd']=='700.00'


@pytest.mark.parametrize('value,expected', [('$9.301,01','9301.01'), ('$0,00','0.00'),
    ('US$ -1.234,50','-1234.50'), ('1,234',None), ('$1,234.50',None), ('=SUM(A1:A2)',None)])
def test_currency_text_is_accepted_only_with_unambiguous_saved_format(value,expected):
    from app.services.sales_bi_parser import number
    result=number(value)
    assert (str(result) if result is not None else None)==expected
