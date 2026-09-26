"""Read saved Excel results. Never evaluate formulas or write to the source workbook."""
from collections import Counter
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO
import json
import re
import unicodedata
from zipfile import ZipFile, BadZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

PARSER_VERSION = 1
MONTHS = ['janeiro', 'fevereiro', 'marco', 'abril', 'maio', 'junho', 'julho',
          'agosto', 'setembro', 'outubro', 'novembro', 'dezembro']
CURRENCIES = {'g$': 'PYG', 'gs': 'PYG', 'r$': 'BRL', 'u$': 'USD', 'us$': 'USD', 'eur': 'EUR', '€': 'EUR'}


class WorkbookError(ValueError):
    pass


def normal(value):
    return ' '.join(''.join(c for c in unicodedata.normalize('NFKD', str(value or '').casefold())
                            if not unicodedata.combining(c)).split())


def number(value):
    # No parsing of formula strings: missing Excel caches remain missing.
    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        return None
    try:
        result = Decimal(str(value))
        return result if result.is_finite() else None
    except InvalidOperation:
        return None


def packed(value):
    return str(value) if value is not None else None


def seller_name(value):
    name = normal(value)
    return {'junior': 'Junior', 'juninho': 'Junior', 'wiss': 'Wissam', 'wissam': 'Wissam'}.get(name, str(value or '').strip().title())


def filename_month(filename):
    words = re.split(r'[^a-z]+', normal(filename))
    return next((i + 1 for i, name in enumerate(MONTHS) if name in words), None)


def parse_workbook(content, *, year=None, month=None, current=False):
    """Only persist sales summaries, coordinates and fingerprints, never customer rows."""
    if len(content) > 20_000_000:
        raise WorkbookError('Planilha maior que o limite de 20 MB.')
    try:
        with ZipFile(BytesIO(content)) as archive:
            if len(archive.infolist()) > 2000 or sum(i.file_size for i in archive.infolist()) > 100_000_000:
                raise WorkbookError('Planilha excede o limite de leitura.')
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=True, keep_links=False)
    except (BadZipFile, KeyError, OSError, ValueError) as exc:
        if isinstance(exc, WorkbookError):
            raise
        raise WorkbookError('Arquivo indisponível ou não é uma planilha XLSX válida.') from None
    try:
        sheets = {}
        for sheet in workbook:
            if normal(sheet.title) not in ('planilha1', 'plhanilha1', 'vendas') and not re.fullmatch(r'semana\s*\d+', normal(sheet.title)):
                continue
            if sheet.max_row and sheet.max_row > 20000 or sheet.max_column and sheet.max_column > 1000:
                raise WorkbookError('Aba de vendas excede o limite de leitura.')
            rows = list(sheet.iter_rows(max_row=min(sheet.max_row or 20000, 20000), max_col=60, values_only=True))
            def get(r, c):
                return rows[r - 1][c - 1] if 0 < r <= len(rows) else None
            total_row = next((r for r in (8, 7) if normal(get(r, 2)) == 'total'), None)
            sellers = {c: seller_name(get(2, c)) for c in range(4, 9) if isinstance(get(2, c), str) and normal(get(2, c))}
            if not total_row or not sellers:
                continue
            totals = {name: number(get(total_row, col)) for col, name in sellers.items()}
            currency_rows = {CURRENCIES[normal(get(r, 3))]: r for r in range(3, 7) if normal(get(r, 3)) in CURRENCIES}
            amounts = {name: {currency: packed(number(get(r, col))) for currency, r in currency_rows.items()} for col, name in sellers.items()}
            raw = [list(r[1:6]) for r in rows[total_row + 1:] if normal(r[1]) in CURRENCIES and number(r[2]) is not None]
            populated = bool(raw) or any(v is not None and v != 0 for v in totals.values())
            fingerprint = sha256(json.dumps(raw, ensure_ascii=False, default=str).encode()).hexdigest() if raw else None
            sheets[sheet.title] = {'rows': rows, 'total': number(get(total_row, 11)), 'sellers': totals, 'currencies': amounts,
                                   'populated': populated, 'fingerprint': fingerprint, 'total_row': total_row}
        mains = [(name, s) for name, s in sheets.items() if normal(name) in ('planilha1', 'plhanilha1', 'vendas')]
        if not mains:
            raise WorkbookError('Não foi encontrado o resumo de vendas neste arquivo.')
        # Some 2023 files have an unrelated blank Planilha1 and sales in Vendas.
        main_name, main = max(mains, key=lambda item: sum(v is not None for row in item[1]['rows'][:15] for v in row[17:]))
        rows = main['rows']
        def get(r, c):
            return rows[r - 1][c - 1] if 0 < r <= len(rows) else None
        def cell(r, c):
            return f'{main_name}!{get_column_letter(c)}{r}'
        published_weeks = []
        total, total_cell = None, None
        for r in range(1, min(16, len(rows) + 1)):
            label = get(r, 18)
            if isinstance(label, str) and re.search(r'\d{1,2}\s*/\s*\d{1,2}', label):
                val = number(get(r, 19))
                pos = (r, 19)
                if val is None:
                    val = number(get(r + 1, 18))
                    pos = (r + 1, 18)
                if val is not None:
                    published_weeks.append({'label': label.strip(), 'total': val, 'cell': cell(*pos)})
            if normal(label).rstrip(': ') in ('total', 'mes', 'total mes'):
                val = number(get(r, 19))
                pos = (r, 19)
                if val is None:
                    val, pos = number(get(r + 1, 18)), (r + 1, 18)
                if val is not None:
                    total, total_cell = val, cell(*pos)
        if month is None:
            dates = [int(m) for w in published_weeks for m in re.findall(r'\d{1,2}\s*/\s*(\d{1,2})', w['label'])]
            month = Counter(m for m in dates if 1 <= m <= 12).most_common(1)[0][0] if dates else None
        if year is None:
            years = {int(v) for row in rows[:16] for v in row[19:22] if number(v) is not None and 2000 <= number(v) <= 2100 and int(v) == v}
            year = next(iter(years)) if len(years) == 1 else None
        if not year or not month:
            raise WorkbookError('Informe o ano e o mês da planilha atual nas fontes do painel.')
        if not current and (total is None or not published_weeks):
            raise WorkbookError('Resumo mensal sem resultado salvo. Salve a planilha no Excel e atualize.')

        # Published week totals are the authority. Match tabs to them; do not add
        # Planilha1 on top of an identical archived week, or guess shifted weeks.
        candidates = [(name, s) for name, s in sheets.items() if s['populated']]
        duplicate_main = any(s['fingerprint'] and s['fingerprint'] == main['fingerprint'] for name, s in candidates if name != main_name)
        if duplicate_main:
            candidates = [(name, s) for name, s in candidates if name != main_name]
        weeks = []
        if current:
            # In the live file the right-hand weekly amounts are typed by hand.
            # Read the saved K8/D8:H8 results instead, so new daily entries appear
            # even before the user updates that monthly recap.
            numbered = [(int(re.search(r'\d+', name).group()), name, s) for name, s in candidates
                        if re.fullmatch(r'semana\s*\d+', normal(name))]
            live = [(name, s) for name, s in candidates if name == main_name]
            if live:
                numbered.append((max((i for i, _, _ in numbered), default=0) + 1, *live[0]))
            for i, name, s in sorted(numbered):
                if s['total'] is None:
                    raise WorkbookError('Resumo semanal sem resultado salvo. Salve a planilha no Excel e atualize.')
                label = published_weeks[i - 1]['label'] if i <= len(published_weeks) else f'Semana {i}'
                weeks.append({'index': i, 'label': label, 'total_usd': packed(s['total']),
                              'source_cell': f"{name}!K{s['total_row']}", 'sheet': name,
                              'sellers': {k: packed(v) for k, v in s['sellers'].items()},
                              'currencies': s['currencies'], 'detail_available': all(v is not None for v in s['sellers'].values())})
            if not weeks:
                raise WorkbookError('A planilha atual ainda não tem semanas preenchidas.')
            total = sum((Decimal(w['total_usd']) for w in weeks), Decimal(0))
            total_cell = None
        else:
            used = set()
            for i, published in enumerate(published_weeks, 1):
                matches = [(name, s) for name, s in candidates if name not in used and s['total'] is not None and abs(s['total'] - published['total']) <= Decimal('0.02')]
                match = matches[0] if len(matches) == 1 else next(((n, s) for n, s in matches if normal(n).replace(' ', '') == f'semana{i}'), None)
                week = {'index': i, 'label': published['label'], 'total_usd': packed(published['total']), 'source_cell': published['cell'],
                        'sellers': {}, 'currencies': {}, 'detail_available': False}
                if match:
                    name, s = match
                    used.add(name)
                    week.update(sellers={k: packed(v) for k, v in s['sellers'].items()}, currencies=s['currencies'],
                                detail_available=all(v is not None for v in s['sellers'].values()), sheet=name)
                weeks.append(week)
        detailed = all(w['detail_available'] for w in weeks)
        names = sorted({name for s in sheets.values() for name in s['sellers']})
        monthly_sellers = {name: {'total_usd': None, 'source_cell': None, 'currencies': {}} for name in names}
        for name in names:
            if detailed:
                values = [number_from_string(w['sellers'].get(name, '0')) for w in weeks]
                monthly_sellers[name]['total_usd'] = packed(sum(values, Decimal(0)))
            currencies = sorted({c for w in weeks for c in w['currencies'].get(name, {})})
            if detailed:
                for c in currencies:
                    vals = [w['currencies'].get(name, {}).get(c, '0') for w in weeks]
                    monthly_sellers[name]['currencies'][c] = packed(sum((number_from_string(v) for v in vals), Decimal(0))) if all(v is not None for v in vals) else None
        # Newer workbooks publish a separate corrected monthly seller result.
        # Consume it as saved, including manual corrections, without exposing formulas.
        for r in range(1, min(12, len(rows) + 1)):
            for c in range(22, 40):
                if current or normal(get(r, c)) != 'total mes':
                    continue
                for col in range(c, min(c + 10, 60)):
                    name = seller_name(get(r + 1, col))
                    value = number(get(r + 2, col))
                    if name in monthly_sellers and value is not None:
                        monthly_sellers[name]['total_usd'] = packed(value)
                        monthly_sellers[name]['source_cell'] = cell(r + 2, col)
        warnings = []
        if not detailed:
            warnings.append('Algumas semanas não têm detalhamento conciliado por vendedor ou moeda; o total mensal e os totais semanais vêm do resumo salvo.')
        if any(v['total_usd'] is None for v in monthly_sellers.values()):
            warnings.append('O ranking por vendedor não inclui este mês quando faltar o resultado do vendedor selecionado.')
        return {'version': PARSER_VERSION, 'year': year, 'month': month, 'current': current, 'total_usd': packed(total),
                'source_cell': total_cell, 'sellers': monthly_sellers, 'weeks': weeks, 'warnings': warnings,
                'currencies_complete': detailed, 'main_sheet': main_name}
    finally:
        workbook.close()


def number_from_string(value):
    return Decimal(value)
